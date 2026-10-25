"""Find candidate Casa 24 words in the private Drive export.

Usage:
    .venv/bin/python scripts/mine_words.py            # reads ./drive_export/
    .venv/bin/python scripts/mine_words.py some/folder

Writes (both git-ignored, never commit them):
    mined_candidates.csv  lang,word,count,in_base_dictionary,offensive,decision
    needs_ocr.txt         PDFs with no extractable text (scanned)

Privacy rules (only the WORD ever leaves this script, never its sentence):
  - emails, URLs, @handles, and any token touching digits are dropped
  - whole lines that mention passwords, accounts, banks, PINs, etc. are dropped
  - ALL-CAPS tokens that are not dictionary words (codes, IDs, acronyms) are dropped
  - words that are not in a dictionary and are usually Capitalized (likely
    names of people, places or brands) are dropped
  - anything in scripts/casa24/blocklist_PRIVATE.txt is dropped
  - files whose path looks like finance, legal or access documents are read
    only to EXCLUDE words: a word that appears only there is never output

Nothing is added to the game. Review the CSV: set `decision` to answer,
guess-only, or reject (fix `lang` to es/en/mixed for unknown words), then run
    .venv/bin/python scripts/merge_words.py mined_candidates.csv
"""

import csv
import os
import re
import sys
import unicodedata
from collections import Counter, defaultdict

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
LANG_DIR = os.path.join(ROOT, "webapp", "data", "languages")
BLOCKLIST = os.path.join(ROOT, "scripts", "casa24", "blocklist_PRIVATE.txt")
TEXT_EXT = {".txt", ".md", ".csv", ".json", ".html", ".htm"}
ALLOWED = set("abcdefghijklmnopqrstuvwxyzñáéíóúü")

# Paths containing any of these are treated as finance / legal / access documents.
SENSITIVE_PATH = re.compile(
    r"financ|finanz|contab|presupuest|budget|factura|invoice|banco|bank|pago|payment|"
    r"legal|contrato|contract|abogad|attorney|lawyer|nda|acuerdo|agreement|"
    r"acceso|access|password|contrase|credencial|credential|login|cuenta|account|clave|"
    r"impuest|tax|nomina|payroll",
    re.IGNORECASE,
)
# Lines containing any of these are dropped completely before tokenizing.
SENSITIVE_LINE = re.compile(
    r"password|contrase|clave|passw|pin\b|usuario|user ?name|login|cuenta|account|"
    r"iban|swift|routing|tarjeta|card|cvv|token|api ?key|secret|2fa|c[oó]digo",
    re.IGNORECASE,
)
STRIP = [
    re.compile(r"\S+@\S+"),                       # emails
    re.compile(r"(https?://|www\.)\S+", re.I),    # URLs
    re.compile(r"[@#]\w+"),                       # handles and hashtags
    re.compile(r"<[^>]+>"),                       # html tags
]

# Small starter lists, flag only (you decide). 5-letter forms matter.
OFFENSIVE = set("""
bitch fucks fucky shits shitty dicks cunts whore pussy penis sluts nigga
twats prick cocks bimbo dildo horny boobs titty rapes raped
putas putos putón verga vergo culos culeo pinga joder polla coños chocha mamón
cagar cagón pajas pajeo zorra perra maric marica cabro chuch pendo
tetas teton tetón follo folla cojón
""".split())


def read_list(lang, name):
    path = os.path.join(LANG_DIR, lang, name.format(lang=lang))
    if not os.path.exists(path):
        return set()
    with open(path, encoding="utf-8") as f:
        return {line.strip() for line in f if line.strip() and not line.startswith("#")}


def load_dictionaries():
    dicts, casa24 = {}, set()
    for lang in ["es", "en"]:
        dicts[lang] = read_list(lang, "{lang}_5words.txt") | read_list(lang, "{lang}_5words_supplement.txt")
        for tier in ["answers", "guesses", "aliases"]:
            casa24 |= read_list(lang, "{lang}_5words_casa24_" + tier + ".txt")
    return dicts, casa24


def load_blocklist():
    if not os.path.exists(BLOCKLIST):
        return set()
    with open(BLOCKLIST, encoding="utf-8") as f:
        return {
            unicodedata.normalize("NFC", line.strip().lower())
            for line in f
            if line.strip() and not line.startswith("#")
        }


def read_file(path, needs_ocr):
    ext = os.path.splitext(path)[1].lower()
    try:
        if ext in TEXT_EXT:
            with open(path, encoding="utf-8", errors="replace") as f:
                return f.read()
        if ext == ".docx":
            import docx

            d = docx.Document(path)
            parts = [p.text for p in d.paragraphs]
            for table in d.tables:
                for row in table.rows:
                    parts += [cell.text for cell in row.cells]
            return "\n".join(parts)
        if ext == ".pdf":
            import pdfplumber

            with pdfplumber.open(path) as pdf:
                text = "\n".join((page.extract_text() or "") for page in pdf.pages)
            if len(text.strip()) < 20:
                needs_ocr.append(path)
                return ""
            return text
    except Exception as e:  # unreadable file: report, keep going
        print(f"  could not read {os.path.basename(path)}: {type(e).__name__}")
    return ""


def tokens(text):
    """Yield (token, original_form) for clean 5-letter letter-only tokens."""
    for line in text.splitlines():
        if SENSITIVE_LINE.search(line):
            continue
        # drop quoted strings (could be passwords or account details)
        line = re.sub(r"[\"“”'‘’`][^\"“”'‘’`]{1,40}[\"“”'‘’`]", " ", line)
        for pattern in STRIP:
            line = pattern.sub(" ", line)
        for raw in re.findall(r"\w+", line):
            raw = unicodedata.normalize("NFC", raw)
            if len(raw) != 5 or not raw.isalpha():
                continue  # wrong length, or touches digits/underscores
            word = raw.lower()
            if all(c in ALLOWED for c in word):
                yield word, raw


def mine(folder):
    dicts, casa24 = load_dictionaries()
    blocked = load_blocklist()
    counts = Counter()          # occurrences in normal documents
    capitalized = Counter()     # how often a word was Capitalized there
    all_caps = Counter()
    seen_sensitive = set()      # words seen in finance/legal/access docs
    needs_ocr, n_files = [], 0

    for dirpath, _, files in os.walk(folder):
        for name in sorted(files):
            path = os.path.join(dirpath, name)
            rel = os.path.relpath(path, folder)
            text = read_file(path, needs_ocr)
            if not text:
                continue
            n_files += 1
            sensitive = bool(SENSITIVE_PATH.search(rel))
            for word, raw in tokens(text):
                if sensitive:
                    seen_sensitive.add(word)
                    continue
                counts[word] += 1
                if raw[0].isupper():
                    capitalized[word] += 1
                if raw.isupper():
                    all_caps[word] += 1

    rows = []
    for word, count in counts.items():
        in_es, in_en = word in dicts["es"], word in dicts["en"]
        known = in_es or in_en
        if word in blocked or word in casa24:
            continue
        if not known and all_caps[word] == count:
            continue  # codes, IDs, acronyms
        if not known and capitalized[word] / count >= 0.6:
            continue  # probably a name of a person, place or brand
        lang = "mixed" if in_es and in_en else "es" if in_es else "en" if in_en else "unknown"
        rows.append({
            "lang": lang,
            "word": word,
            "count": count,
            "in_base_dictionary": "yes" if known else "no",
            "offensive": "yes" if word in OFFENSIVE else "",
            "decision": "",
        })
    rows.sort(key=lambda r: (-r["count"], r["word"]))
    return rows, needs_ocr, n_files, len(seen_sensitive - set(counts))


def main(folder):
    if not os.path.isdir(folder):
        sys.exit(f"Folder not found: {folder}. Put the Drive export in ./drive_export/ first.")
    if not os.path.exists(BLOCKLIST):
        print("Note: no scripts/casa24/blocklist_PRIVATE.txt yet (see blocklist.example.txt).")
    rows, needs_ocr, n_files, n_sensitive_only = mine(folder)
    with open(os.path.join(ROOT, "mined_candidates.csv"), "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=["lang", "word", "count", "in_base_dictionary", "offensive", "decision"])
        writer.writeheader()
        writer.writerows(rows)
    with open(os.path.join(ROOT, "needs_ocr.txt"), "w", encoding="utf-8") as f:
        f.write("\n".join(os.path.relpath(p, folder) for p in needs_ocr) + ("\n" if needs_ocr else ""))
    by_lang = Counter(r["lang"] for r in rows)
    print(f"Read {n_files} files. {len(rows)} candidate words: {dict(by_lang)}")
    print(f"Flagged offensive: {sum(1 for r in rows if r['offensive'])}")
    print(f"Skipped {n_sensitive_only} words seen only in finance/legal/access files")
    print(f"PDFs needing OCR: {len(needs_ocr)} (see needs_ocr.txt)")
    print("Wrote mined_candidates.csv. Nothing was added to the game.")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, "drive_export"))
